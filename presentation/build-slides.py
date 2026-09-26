"""Build presentation/cce-2026-slides.pptx from cengage-template.pptx.

Source of the content: presentation/slides-outline.md (revision 3).
Rules: one placeholder per column, multi-paragraph text inside it; no extra
text boxes except a title on the blank Image layout, a one-line caption or
bottom band where the outline calls for one, and the five-part footer strip.

Run:  python presentation/build-slides.py
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "presentation" / "cengage-template.pptx"
OUT = ROOT / "presentation" / "cce-2026-slides.pptx"
IMG = ROOT / "assets" / "images"

# --------------------------------------------------------------------------- helpers

def layout(prs, name):
    for m in prs.slide_masters:
        for l in m.slide_layouts:
            if l.name == name:
                return l
    raise KeyError(name)


def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    raise KeyError(f"placeholder idx {idx} not on slide")


def add_runs(paragraph, text, size=None, color=None):
    """Write text into a paragraph; **bold** segments become bold runs."""
    parts = re.split(r"\*\*(.+?)\*\*", text)
    for i, part in enumerate(parts):
        if not part:
            continue
        r = paragraph.add_run()
        r.text = part
        if i % 2 == 1:
            r.font.bold = True
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color


def fill(shape, lines, size=None, align=None, anchor=None):
    """Fill a placeholder or text box with one paragraph per line."""
    tf = shape.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if align is not None:
            p.alignment = align
        add_runs(p, line, size=size)


def textbox(slide, left, top, width, height, lines, size=16, bold=False,
            align=None, anchor=MSO_ANCHOR.TOP, fill_rgb=None):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if fill_rgb is not None:
        tb.fill.solid()
        tb.fill.fore_color.rgb = fill_rgb
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.15)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        add_runs(p, line, size=size)
        if bold:
            for r in p.runs:
                r.font.bold = True
    return tb


def remove(shape):
    el = shape._element
    el.getparent().remove(el)


def picture_in(slide, shape_or_rect, path, pad=0.0):
    """Place a picture fitted inside a placeholder's box (placeholder removed)
    or inside an (left, top, width, height) rectangle in inches."""
    if isinstance(shape_or_rect, tuple):
        l, t, w, h = shape_or_rect
    else:
        l, t, w, h = (shape_or_rect.left / 914400, shape_or_rect.top / 914400,
                      shape_or_rect.width / 914400, shape_or_rect.height / 914400)
        remove(shape_or_rect)
    l, t, w, h = l + pad, t + pad, w - 2 * pad, h - 2 * pad
    pic = slide.shapes.add_picture(str(IMG / path), Inches(l), Inches(t))
    # fit, preserving aspect ratio, then centre
    ratio = min(Inches(w) / pic.width, Inches(h) / pic.height)
    pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
    pic.left = int(Inches(l) + (Inches(w) - pic.width) / 2)
    pic.top = int(Inches(t) + (Inches(h) - pic.height) / 2)
    return pic


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def clean(slide):
    """Drop any placeholder left empty so no 'click to add' boxes remain."""
    for p in list(slide.placeholders):
        if p.placeholder_format.idx == 0:
            continue
        if p.placeholder_format.type is not None and "PICTURE" in str(p.placeholder_format.type):
            remove(p)
        elif p.has_text_frame and not p.text_frame.text.strip():
            remove(p)


# --------------------------------------------------------------------------- build

prs = Presentation(TEMPLATE)
L = lambda n: layout(prs, n)

original = list(prs.slides._sldIdLst)  # Event Title, Housekeeping, 5 samples, Survey, Questions, Thank you
new_ids = []


def new(layout_name, title=None):
    s = prs.slides.add_slide(L(layout_name))
    new_ids.append(prs.slides._sldIdLst[-1])
    if title is not None and s.shapes.title is not None:
        s.shapes.title.text = title
    return s


# 1 -------------------------------------------------------------------------
s = new("Session Title", "AI Agents in VS Code for Academic Work")
fill(ph(s, 10), ["From Prompting to Workflows  ·  Nim Dvir, PhD, MBA  ·  University at Albany, Massry School of Business"])
picture_in(s, (8.0, 1.0, 3.2, 3.2), "headshot2025-cloud.jpg")
textbox(s, 8.0, 4.4, 3.2, 1.2,
        ["**Scan to follow along**", "QR code and short URL: TODO (add qr-site.png once the site URL is final)"],
        size=14, align=PP_ALIGN.CENTER)
notes(s, """
Image: headshot2025-cloud.jpg, qr-site.png (TODO once the URL is final). QR large enough to scan from the back row; this slide is up before the session starts.
First words: "This website is the presentation — and AI agents helped me build it. Let's change it."
TODO: date, time, room.
""")

# 2 -------------------------------------------------------------------------
s = new("1 column with footer", "Let's change it")
fill(ph(s, 18), ["Watch your phone."])
picture_in(s, ph(s, 24), "github-pages-publishing-source.png", pad=0.15)
notes(s, """
Image: capture of the live Intro page with the red-dot banner highlighted (TODO after Pages is live). Placeholder for now: github-pages-publishing-source.png.
Switch to VS Code. Agent mode. Type the saved prompt (/update-live-banner):
"In docs/content/intro.md, find the line marked LIVE BANNER and replace the text after the red dot with a short greeting to the CCE 2026 audience. Keep it one line. Rebuild the site with python docs/build.py, then commit with the message "Update live banner" and push to main."
While it runs, narrate only what is visible: it opened a file, changed one line, ran a script, saved a checkpoint, published. Refresh.
Do not explain Git or Pages here. Say: "How that works is slide 13. First, why you would want it."
Fallback: pre-pushed commit to cherry-pick. Reset before the session: git checkout demo-start -- live/ docs/content/intro.md && git clean -fd live/
""")

# 3 -------------------------------------------------------------------------
s = new("3 column with footer", "Every semester, the same work")
fill(ph(s, 24), ["📅  Dates"]); fill(ph(s, 22), ["📊  Spreadsheets"]); fill(ph(s, 23), ["📝  Piles"])
fill(ph(s, 26), ["Update the syllabus, the schedule, the policies, the assignments.", "Again."])
fill(ph(s, 27), ["Course feedback you meant to analyze."])
fill(ph(s, 28), ["The same rubric, applied a hundred times."])
textbox(s, 0.42, 5.2, 12.49, 1.0, ["What if the AI could open your files?"], size=28, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, fill_rgb=RGBColor(0xFF, 0xFF, 0xFF))
notes(s, """
Image: none; the three headings carry the icons. No generated art.
Ask the room which one they did last week. This is the frame for every demo: the three columns are Demos 1, 2, and 3.
The promise, in one breath: you have seen it work. In the next 40 minutes: what is thinking (Brain), where it works (Infrastructure), what you tell it (Task), how to work with it reliably (Workflow), and three of your own tasks (Examples). That sentence is the map; the footer strip carries it from here on.
""")

# 4 -------------------------------------------------------------------------
s = new("2 column with footer", "An agent doesn't answer. It works.")
remove(ph(s, 18)); remove(ph(s, 24))
picture_in(s, (0.42, 1.17, 5.84, 5.14), "generated/1-what-makes-an-agent.png")
fill(ph(s, 20), ["What is an AI agent?"])
fill(ph(s, 26), [
    "An AI agent uses a model, context, instructions, and tools to work toward a goal through multiple steps.",
    "",
    "**Chat:** you ask → it answers → you do the work",
    "**Agent:** you give a goal → it inspects → plans → acts → checks",
    "",
    "**Chat is the interface. An agent is a way the AI can operate.**",
])
notes(s, """
Image: generated/1-what-makes-an-agent.png.
Everyone here uses chat already; don't talk it down. The shift is who carries the output and coordinates the steps. On slide 2, nobody copy-pasted.
What it cannot be assumed to do: know facts it was not given, resolve ambiguity every time, judge what matters pedagogically, replace your review. That sets up slide 14.
Today the agent runs locally, next to your files, through GitHub Copilot in VS Code. Other places to run one are on the site.
""")

# 5 -------------------------------------------------------------------------
s = new("2 column with footer", "It can only see what you give it")
remove(ph(s, 18)); remove(ph(s, 24))
picture_in(s, (0.42, 1.17, 5.84, 5.14), "generated/2-tokens-and-the-context-window.png")
fill(ph(s, 20), ["Context, model, tokens"])
fill(ph(s, 26), [
    "**Context** = prompt + conversation + files + instructions + tool results",
    "",
    "**Model** = the engine",
    "**Tokens** = the units it reads and writes; everything counts",
    "",
    "**Give it the material relevant to the task, not everything you have.**",
])
notes(s, """
Image: generated/2-tokens-and-the-context-window.png. Optional small inset: copilot-model-picker.png.
Do not teach token arithmetic; capacity varies by model. The lesson is relevance.
Second line to say: "Don't keep explaining the project. Put the project where the agent can understand it." That is why the demos run inside a project folder, and it pays off on slide 12 (AGENTS.md).
If asked about VS Code's own vocabulary (harness, model, context, tools, prompt): the mapping table is on the Explanation page.
""")

# 6 -------------------------------------------------------------------------
s = new("1 column with footer", "VS Code: your files, with an AI beside them")
fill(ph(s, 18), ["VS Code + your files + Copilot = an agentic workspace   ·   Setup takes about twenty minutes. Guide on the site."])
picture_in(s, ph(s, 24), "vscode-interface.png", pad=0.1)
notes(s, """
Image: vscode-interface.png. Add three callouts in PowerPoint if wanted: your files, the editor, Copilot Chat. Optional inset: vscode-extensions.png.
A text editor edits plain-text files directly (.md, .csv, .html, .py). VS Code grew into a whole workspace: files, editor, terminal, extensions, AI. You do not need to become a software engineer.
When practical, keep working source material in plain text; export to Word or PDF when needed. Word and PDF can work with the right tooling; plain text is just more transparent to the agent. (generated/3-plain-text-vs-complex-file-formats.png is on the site if someone asks.)
Extensions are apps for VS Code; Copilot is one. Verified students can access Copilot Student at no cost, and eligible verified teachers can receive Copilot Pro at no cost through GitHub Education. Check current eligibility before the session.
Setup steps (site, not slide): GitHub account · GitHub Education · install VS Code · sign in · clone a repository · enable Copilot · extensions. TODO: Education URL.
""")

# 7 -------------------------------------------------------------------------
s = new("3 column with footer", "Goal + Context + Constraints")
fill(ph(s, 24), ["Goal"]); fill(ph(s, 22), ["Context"]); fill(ph(s, 23), ["Constraints"])
fill(ph(s, 26), ["✗  \"Fix my course.\"", "", "✓  \"Update this syllabus and schedule for Spring 2027.\""])
fill(ph(s, 27), ["The files.", "", "Syllabus, schedule, policies, a CSV, the existing site."])
fill(ph(s, 28), ["Don't change grading weights.", "Don't invent dates.", "Flag uncertainty.", "Don't commit."])
textbox(s, 0.42, 4.8, 12.49, 1.45,
        ["The course files are in live/syllabus/.",
         "Update the course for Spring 2027 without changing grading weights. Identify inconsistent dates and flag anything requiring judgment."],
        size=16, anchor=MSO_ANCHOR.MIDDLE, fill_rgb=RGBColor(0xF2, 0xF2, 0xF2))
notes(s, """
Image: none; the three columns are the graphic. Optional, if a picture is wanted: "Three blocks labeled GOAL, CONTEXT, CONSTRAINTS joined by plus signs, an equals sign, one block labeled TASK; CONTEXT holds document and table icons, CONSTRAINTS a padlock; flat, white background, one accent color."
Walk the boxed prompt: line 1 is context; sentence 2 is goal plus one constraint; sentence 3 is two more. Three lines. That is the exact prompt for the next demo.
Say, don't show: name the authoritative source, tell it to inspect before editing, ask it to verify afterward. Move knowledge from the prompt into the workspace.
""")

# 8 -------------------------------------------------------------------------
s = new("4 column with footer", "Ask. Plan. Agent. Then watch the diff.")
fill(ph(s, 26), ["Ask"]); fill(ph(s, 29), ["Plan"]); fill(ph(s, 30), ["Agent"]); fill(ph(s, 32), ["Review"])
fill(ph(s, 28), ["Understand.", "No edits."])
fill(ph(s, 33), ["Propose the changes.", "Waits for you."])
fill(ph(s, 34), ["Does the work.", "Runs tools."])
fill(ph(s, 35), ["You.", "Files changed, the diff, Git status.", "",
                 "**\"Completed\" does not necessarily mean \"correct.\"**"])
picture_in(s, (0.42, 3.4, 2.85, 1.6), "copilot-ask-plan-agent.png")
notes(s, """
Image: copilot-ask-plan-agent.png (the mode picker). Optional: vscode-diff-editor.png in the Review column.
Same chat panel, three modes; the fourth column is you. Plan is the one faculty underuse. With manual permissions, Copilot requests approval for actions that are not already allowed; review commands before approving them.
Git makes the loop safe: commit before, review the diff after, revert if needed. Visible, reviewable, reversible.
The line to land: "In the next demo, watch the diff, not the chat." That is where supervision happens.
""")

# 9 -------------------------------------------------------------------------
def demo_card(title, intro, heading, watch, picture, part, note):
    s = new("Text with intro text and footer", title)
    fill(ph(s, 21), [intro])
    fill(ph(s, 18), [heading])
    fill(ph(s, 24), [watch])
    picture_in(s, (0.42, 2.95, 12.49, 3.4), picture)
    notes(s, note)
    return s

demo_card(
    "Your syllabus, updated while you watch",
    "4 course files with inconsistent dates  →  agent  →  updated files + a list of flagged decisions",
    "/update-syllabus   ·   Plan → Agent → diff",
    "**Watch for:** reads all four files · proposes first · flags what needs judgment",
    "vscode-diff-editor.png", 4, """
Image: real screenshot of the multi-file diff from rehearsal (TODO, plan-claude.md step 13). Until then vscode-diff-editor.png.
Run Plan first and read the plan aloud. Then Agent. Then open the diff and stay there.
Fictional course, deliberately inconsistent dates. Fallback: live/generated/expected/syllabus/. Six minutes.
Aside if time: the same four files can become a six-page course website with one more prompt (/build-course-site). Nothing on those pages is new information. Otherwise leave it for slide 13.
""")

# 10 ------------------------------------------------------------------------
demo_card(
    "From a feedback spreadsheet to one chart and a paragraph",
    "student-feedback.csv  →  inspect  →  analyze  →  visualize  →  interpret",
    "/analyze-feedback   ·   Agent",
    "**Watch for:** looks at the columns first · runs Python · one chart · \"Do not invent data.\"",
    "copilot-agent-browser-validation.png", 4, """
Image: the real chart from live/generated/expected/data/ plus a CSV screenshot (TODO). Until then copilot-agent-browser-validation.png.
Point at the moment it inspects the columns before analyzing: that is "inspect" from slide 4 happening for real.
Synthetic data, about 200 rows, a few messy cells on purpose. Say so; they will ask. Fallback: expected output folder. Five minutes.
""")

# 11 ------------------------------------------------------------------------
s = new("1 column with footer", "One submission is a prompt. A hundred is a workflow.")
fill(ph(s, 18), ["assignment + rubric + submission  →  apply rubric  →  evidence  →  score  →  draft feedback  →  instructor review"])
picture_in(s, ph(s, 24), "generated/6-ai-assisted-grading-workflow.png", pad=0.1)
notes(s, """
Image: generated/6-ai-assisted-grading-workflow.png as the poster frame; replace with the 20–30 second grading clip (TODO, plan-claude.md step 14).
Three synthetic submissions. The skill standardizes the grading procedure, rubric application, and output format across submissions. The instructor reviews, edits, and approves every result.
Emphasize the skill, not "AI grades students." This is the third column from slide 3.
""")

# 12 ------------------------------------------------------------------------
s = new("2 column with footer", "From prompting to workflows")
fill(ph(s, 18), ["The subtitle of this talk"])
fill(ph(s, 24), [
    "**PROMPT** — one task, typed once",
    "**REUSABLE PROMPT** — .github/prompts/update-syllabus.prompt.md",
    "**PROJECT INSTRUCTIONS** — AGENTS.md: \"Always work this way in this project.\"",
    "**SKILL** — .github/skills/grade-submission/SKILL.md",
    "**WORKFLOW** — several steps, chained and repeatable",
])
picture_in(s, ph(s, 26), "generated/4-from-one-prompt-to-a-workflow.png")
notes(s, """
Image: generated/4-from-one-prompt-to-a-workflow.png. Optional: vscode-create-instructions.png.
Every rung is a real file in this repository; open one live if there is a minute. Type /update-syllabus next semester, change one year, run again.
Prompt = "Do this now." AGENTS.md = "Always work this way in this project." This repo's rules: don't invent facts, inspect before editing, don't commit unless asked, synthetic data only.
Different agents support different instruction formats. AGENTS.md works across several agent environments; others also have native formats such as CLAUDE.md.
""")

# 13 ------------------------------------------------------------------------
s = new("2 column with footer", "How I built this website")
fill(ph(s, 18), ["The plumbing behind slide 2"])
fill(ph(s, 24), [
    "outlines + prompts + template → agent → Markdown → build script → GitHub → GitHub Pages → your phone",
    "",
    "**Clone · edit · commit · push.**",
    "Git keeps the history. GitHub keeps the copy. Pages publishes the docs/ folder.",
    "",
    "**Nothing here needed a web developer.**",
])
picture_in(s, ph(s, 26), "github-pages-publishing-source.png")
notes(s, """
Image: github-pages-publishing-source.png (Pages from main / docs). Optional insets: vscode-clone-repository.png, vscode-source-control-graph.png.
This is the plumbing behind slide 2, explained only now that they have seen it pay off twice. Four verbs, no Git lesson. A repository is a project folder whose files and history are tracked with Git; when an agent can change many files, that history is what makes it safe.
The site's four pages are plain Markdown in docs/content/; a short script renders them. The whole repo is on the site: clone it and start from it.
Aside, cut first: same trick turns course files into a course site (generated/5-from-course-files-to-an-interactive-website.png).
Other brains, in one sentence: Copilot's model picker offers GPT, Claude, Gemini, and others; Claude and Codex also run in VS Code. The workflow does not depend on which one you pick. Table on the site. Never publish an API key in a repository.
""")

# 14 ------------------------------------------------------------------------
s = new("2 column with footer", "Delegate work, not responsibility")
fill(ph(s, 18), ["Four words"])
fill(ph(s, 24), [
    "**Protect** sensitive data.",
    "**Constrain** what may change.",
    "**Verify** the result.",
    "**Own** the decision.",
])
fill(ph(s, 20), ["What to try first"])
fill(ph(s, 26), [
    "1.  Open a real project folder in VS Code",
    "2.  Ask Copilot to inspect it",
    "3.  Give it one multi-file task",
    "4.  Create an AGENTS.md",
    "5.  Turn one repetitive task into a workflow",
])
notes(s, """
Image: none on the slide; add four native icons (shield, padlock, magnifier, signature) if wanted. No generated art.
Protect: sensitive student or research data requires institutionally approved tools, storage, and handling; use synthetic data for experimentation and demonstrations. Everything shown today was synthetic.
Verify: agents can invent plausible information, misread ambiguous requirements, and make technically valid but undesirable changes. Own: be explicit with students about how AI is used; the grade, the policy, and the decision remain yours.
Steps 1 and 2 change no files. Start with one folder and one annoying repeated task.
""")

# 15 ------------------------------------------------------------------------
s = new("2 column with footer", "Take it with you")
fill(ph(s, 18), ["On the site"])
fill(ph(s, 24), [
    "These slides · setup guide · every prompt · the demo files · the repository · recordings",
    "",
    "github.com/nimdvir/cce-2026",
    "",
    "Nim Dvir · nimdvir.com · linkedin.com/in/nimdvir",
])
fill(ph(s, 20), ["Scan once more"])
fill(ph(s, 26), ["QR code and short URL: TODO (add qr-site.png once the site URL is final)"])
picture_in(s, (7.08, 2.5, 5.84, 3.8), "headshot2025-cloud.jpg")
notes(s, """
Image: qr-site.png (TODO), same code as slide 1; headshot2025-cloud.jpg small.
The site is the durable artifact; the deck is thin on purpose.
""")

# closing template slides: notes only ---------------------------------------
for sid in original:
    pass

# --------------------------------------------------------------------------- order and cleanup
sldIdLst = prs.slides._sldIdLst
for el in new_ids:  # only the slides built here; the template's own slides stay untouched
    clean(prs.slides.get(int(el.get("id"))))

ev, hk, *samples, survey, questions, thanks = original
for el in list(sldIdLst):
    sldIdLst.remove(el)
for el in [ev, hk, *new_ids, survey, questions, thanks]:
    sldIdLst.append(el)
for el in samples:
    prs.part.drop_rel(el.rId)

# notes on the template's own closing slides
slides = list(prs.slides)
notes(slides[-2], "Keep the site open in case someone asks to revisit a demo. Footer: nimdvir.com · github.com/nimdvir/cce-2026")
notes(slides[-1], 'Add headshot and QR if the layout allows. Close with: "Don\'t just ask AI questions. Give it a workspace, context, rules, and a goal."')

prs.save(OUT)
print(f"wrote {OUT.relative_to(ROOT)} with {len(prs.slides)} slides")
